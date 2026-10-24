"""Slide composition preserves transparency in the embedded image."""

import json
import runpy
from io import BytesIO
from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation

SCRIPT = (
    Path(__file__).resolve().parents[2]
    / "skills/public/ppt-generation/scripts/generate.py"
)


@pytest.mark.parametrize("mode", ["RGBA", "LA", "P", "RGB", "L"])
def test_ppt_preserves_png_transparency(tmp_path, mode):
    source = tmp_path / "slide.png"
    if mode == "P":
        image = Image.new("P", (64, 64), 0)
        image.putpalette([0, 0, 0, 240, 80, 40] + [0] * 762)
        image.putpixel((48, 32), 1)
        image.save(source, transparency=0)
    elif mode in {"RGB", "L"}:
        image = Image.new(mode, (64, 64), 0)
        image.putpixel((48, 32), (240, 80, 40) if mode == "RGB" else 160)
        image.save(source, transparency=(0, 0, 0) if mode == "RGB" else 0)
    else:
        image = Image.new(mode, (64, 64), (0, 0, 0, 0) if mode == "RGBA" else (0, 0))
        image.putpixel((48, 32), (240, 80, 40, 128) if mode == "RGBA" else (160, 128))
        image.save(source)
    image.close()
    with Image.open(source) as original:
        expected = original.convert("RGBA")

    plan = tmp_path / "plan.json"
    plan.write_text(
        json.dumps({"slides": [{"title": "Transparent slide"}]}), encoding="utf-8"
    )
    output = tmp_path / "slides.pptx"
    generate_ppt = runpy.run_path(str(SCRIPT))["generate_ppt"]

    generate_ppt(str(plan), [str(source)], str(output))

    presentation = Presentation(output)
    picture = presentation.slides[0].shapes[0]
    with Image.open(BytesIO(picture.image.blob)) as embedded:
        assert embedded.format == "PNG"
        assert embedded.convert("RGBA").tobytes() == expected.tobytes()
    assert (
        "Title: Transparent slide"
        in presentation.slides[0].notes_slide.notes_text_frame.text
    )


@pytest.mark.parametrize("mode", ["RGB", "L", "P"])
def test_ppt_keeps_opaque_images_as_jpeg(tmp_path, mode):
    source = tmp_path / "slide.png"
    with Image.new(mode, (64, 64), 100) as image:
        image.save(source)
    plan = tmp_path / "plan.json"
    plan.write_text("{}", encoding="utf-8")
    output = tmp_path / "slides.pptx"

    runpy.run_path(str(SCRIPT))["generate_ppt"](str(plan), [str(source)], str(output))

    picture = Presentation(output).slides[0].shapes[0]
    with Image.open(BytesIO(picture.image.blob)) as embedded:
        assert embedded.format == "JPEG"
        assert embedded.size == (64, 64)
