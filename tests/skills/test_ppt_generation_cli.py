"""The ppt-generation CLI must expose composition failures to its shell caller.

These tests execute the real script entry point. Image fixtures are tiny
synthetic files; no image API or Docker is involved.
"""

import json
import runpy
import subprocess
import sys
from io import BytesIO
from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation

SCRIPT = (
    Path(__file__).resolve().parents[2]
    / "skills/public/ppt-generation/scripts/generate.py"
)


def _write_image(
    path: Path, size: tuple[int, int], color: tuple[int, int, int]
) -> None:
    with Image.new("RGB", size, color) as image:
        image.save(path, format="JPEG")


def _run(
    plan: Path, images: list[Path], output: Path
) -> subprocess.CompletedProcess[str]:
    return subprocess.run(
        [
            sys.executable,
            str(SCRIPT),
            "--plan-file",
            str(plan),
            "--slide-images",
            *[str(image) for image in images],
            "--output-file",
            str(output),
        ],
        capture_output=True,
        text=True,
        check=False,
        timeout=30,
    )


def test_missing_slide_image_raises(tmp_path):
    plan = tmp_path / "plan.json"
    plan.write_text("{}", encoding="utf-8")
    generate_ppt = runpy.run_path(str(SCRIPT))["generate_ppt"]

    with pytest.raises(FileNotFoundError, match="Slide image not found"):
        generate_ppt(
            str(plan), [str(tmp_path / "missing.jpg")], str(tmp_path / "out.pptx")
        )


def test_cli_missing_slide_image_exits_nonzero(tmp_path):
    plan = tmp_path / "plan.json"
    plan.write_text(
        json.dumps({"slides": [{"title": "One"}, {"title": "Two"}]}),
        encoding="utf-8",
    )
    present = tmp_path / "slide-01.jpg"
    _write_image(present, (32, 18), (12, 34, 56))
    missing = tmp_path / "slide-02.jpg"
    output = tmp_path / "presentation.pptx"
    output.write_bytes(b"previous pptx")

    result = _run(plan, [present, missing], output)

    assert result.returncode == 1
    assert "Slide image not found" in result.stderr
    assert str(missing) in result.stderr
    assert result.stdout == ""
    assert output.read_bytes() == b"previous pptx"


@pytest.mark.parametrize(
    ("content", "diagnostic"),
    [
        (None, "plan.json"),
        ("{", "Expecting property name"),
        ("[]", "has no attribute 'get'"),
        ('{"slides": ["not-a-slide"]}', "has no attribute 'get'"),
    ],
    ids=["missing-plan", "invalid-json", "plan-not-object", "slide-metadata-mismatch"],
)
def test_cli_plan_failure_exits_nonzero(tmp_path, content, diagnostic):
    plan = tmp_path / "plan.json"
    if content is not None:
        plan.write_text(content, encoding="utf-8")
    image = tmp_path / "slide.jpg"
    _write_image(image, (32, 18), (1, 2, 3))
    output = tmp_path / "presentation.pptx"
    output.write_bytes(b"previous pptx")

    result = _run(plan, [image], output)

    assert result.returncode == 1
    assert "Error while generating presentation:" in result.stderr
    assert diagnostic in result.stderr
    assert result.stdout == ""
    assert output.read_bytes() == b"previous pptx"


def test_cli_non_image_file_exits_nonzero(tmp_path):
    plan = tmp_path / "plan.json"
    plan.write_text("{}", encoding="utf-8")
    image = tmp_path / "slide.jpg"
    image.write_bytes(b"this is not an image")
    output = tmp_path / "presentation.pptx"
    output.write_bytes(b"previous pptx")

    result = _run(plan, [image], output)

    assert result.returncode == 1
    assert "Error while generating presentation:" in result.stderr
    assert "cannot identify image file" in result.stderr
    assert result.stdout == ""
    assert output.read_bytes() == b"previous pptx"


def test_cli_success_prints_stdout_and_writes_pptx(tmp_path):
    first = tmp_path / "slide-01.jpg"
    second = tmp_path / "slide-02.jpg"
    _write_image(first, (32, 18), (10, 20, 30))
    _write_image(second, (16, 16), (200, 10, 10))
    plan = tmp_path / "plan.json"
    plan.write_text(
        json.dumps(
            {
                "title": "Tiny",
                "aspect_ratio": "16:9",
                "slides": [
                    {"title": "One", "subtitle": "Hello", "key_points": ["Alpha"]},
                    {"title": "Two"},
                ],
            }
        ),
        encoding="utf-8",
    )
    output = tmp_path / "presentation.pptx"

    result = _run(plan, [first, second], output)

    assert result.returncode == 0
    assert result.stderr == ""
    assert "Successfully generated presentation with 2 slides" in result.stdout
    assert str(output) in result.stdout
    presentation = Presentation(output)
    assert len(presentation.slides) == 2
    assert "Title: One" in presentation.slides[0].notes_slide.notes_text_frame.text
    assert "Title: Two" in presentation.slides[1].notes_slide.notes_text_frame.text
    with Image.open(BytesIO(presentation.slides[0].shapes[0].image.blob)) as embedded:
        assert embedded.format == "JPEG"
        assert embedded.size == (32, 18)
