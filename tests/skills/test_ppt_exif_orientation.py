import json
import runpy
from io import BytesIO
from pathlib import Path

import pytest
from PIL import Image, ImageDraw
from pptx import Presentation

SCRIPT = (
    Path(__file__).resolve().parents[2]
    / "skills/public/ppt-generation/scripts/generate.py"
)
generate_ppt = runpy.run_path(str(SCRIPT))["generate_ppt"]
COLORS = {"R": (240, 0, 0), "G": (0, 220, 0), "B": (0, 0, 240), "Y": (240, 220, 0)}


@pytest.mark.parametrize("aspect_ratio", ["16:9", "4:3"])
@pytest.mark.parametrize(
    ("orientation", "corners"),
    [
        (0, "RGBY"),
        (1, "RGBY"),
        (2, "GRYB"),
        (3, "YBGR"),
        (4, "BYRG"),
        (5, "RBGY"),
        (6, "BRYG"),
        (7, "YGBR"),
        (8, "GYRB"),
    ],
)
def test_ppt_preserves_display_orientation(
    tmp_path, aspect_ratio, orientation, corners
):
    image_path = tmp_path / "photo.jpg"
    with Image.new("RGB", (96, 64)) as image:
        draw = ImageDraw.Draw(image)
        for box, color in zip(
            [(0, 0, 47, 31), (48, 0, 95, 31), (0, 32, 47, 63), (48, 32, 95, 63)],
            "RGBY",
            strict=True,
        ):
            draw.rectangle(box, fill=COLORS[color])
        exif = image.getexif()
        if orientation:
            exif[274] = orientation
        image.save(image_path, quality=100, subsampling=0, exif=exif)
    plan_path = tmp_path / "plan.json"
    plan_path.write_text(
        json.dumps({"aspect_ratio": aspect_ratio, "slides": [{"title": "Photo"}]}),
        encoding="utf-8",
    )
    output_path = tmp_path / "presentation.pptx"

    generate_ppt(str(plan_path), [str(image_path)], str(output_path))

    presentation = Presentation(output_path)
    picture = presentation.slides[0].shapes[0]
    expected_size = (64, 96) if orientation >= 5 else (96, 64)
    with Image.open(BytesIO(picture.image.blob)) as embedded:
        assert embedded.size == expected_size
        assert 274 not in embedded.getexif()
        width, height = embedded.size
        points = [
            (width // 4, height // 4),
            (3 * width // 4, height // 4),
            (width // 4, 3 * height // 4),
            (3 * width // 4, 3 * height // 4),
        ]
        for point, color in zip(points, corners, strict=True):
            assert embedded.getpixel(point) == pytest.approx(COLORS[color], abs=10)
    assert picture.width / picture.height == pytest.approx(
        expected_size[0] / expected_size[1], rel=1e-5
    )
    assert "Title: Photo" in presentation.slides[0].notes_slide.notes_text_frame.text
