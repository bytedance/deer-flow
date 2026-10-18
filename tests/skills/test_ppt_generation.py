import json
import runpy
from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation

SCRIPT = Path(__file__).resolve().parents[2] / "skills/public/ppt-generation/scripts/generate.py"
generate_ppt = runpy.run_path(str(SCRIPT))["generate_ppt"]


@pytest.mark.parametrize("aspect_ratio", ["16:9", "4:3"])
@pytest.mark.parametrize("image_size", [(320, 80), (80, 320), (120, 120), (160, 120)])
def test_slide_images_are_centered_without_cropping(tmp_path, aspect_ratio, image_size):
    plan_path = tmp_path / "plan.json"
    plan_path.write_text(
        json.dumps({"aspect_ratio": aspect_ratio, "slides": [{"title": "Centered image"}]}),
        encoding="utf-8",
    )
    image_path = tmp_path / "slide.png"
    Image.new("RGB", image_size, "navy").save(image_path)
    output_path = tmp_path / "presentation.pptx"

    generate_ppt(str(plan_path), [str(image_path)], str(output_path))

    presentation = Presentation(output_path)
    assert len(presentation.slides) == 1
    slide = presentation.slides[0]
    assert len(slide.shapes) == 1
    picture = slide.shapes[0]
    right = presentation.slide_width - picture.left - picture.width
    bottom = presentation.slide_height - picture.top - picture.height

    assert min(picture.left, picture.top, right, bottom) >= 0
    # Allow two EMUs for the existing inch-to-integer coordinate rounding.
    assert abs(picture.left - right) <= 2
    assert abs(picture.top - bottom) <= 2
    assert picture.width / picture.height == pytest.approx(image_size[0] / image_size[1], rel=1e-5)
    assert picture.width == presentation.slide_width or picture.height == presentation.slide_height
    assert picture.image.size == image_size
    assert "Title: Centered image" in slide.notes_slide.notes_text_frame.text
