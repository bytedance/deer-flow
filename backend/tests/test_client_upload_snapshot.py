"""Embedded upload and conversion must consume the same captured bytes."""

import asyncio
import getpass
import os
import shutil
import stat
import subprocess
import sys
from io import BytesIO
from pathlib import Path

import pytest

from deerflow.client import DeerFlowClient
from deerflow.uploads.companions import resolve_companion


@pytest.fixture
def upload_case(monkeypatch, tmp_path):
    import deerflow.client as client_module

    uploads = tmp_path / "thread/user-data/uploads"
    uploads.mkdir(parents=True)
    monkeypatch.setattr(client_module, "ensure_uploads_dir", lambda *args, **kwargs: uploads)
    # upload_files has no dependency on initialized instance fields.
    return DeerFlowClient.__new__(DeerFlowClient), uploads


def call_upload(client, source, in_loop):
    if not in_loop:
        return client.upload_files("snapshot-thread", [source])

    async def call():
        return client.upload_files("snapshot-thread", [source])

    return asyncio.run(call())


@pytest.mark.parametrize("in_loop", [False, True])
@pytest.mark.parametrize("change", ["edit", "remove"])
@pytest.mark.parametrize("existing", [False, True])
def test_conversion_keeps_uploaded_version(upload_case, monkeypatch, tmp_path, in_loop, change, existing):
    import deerflow.client as client_module
    import deerflow.utils.file_conversion as conversion

    client, uploads = upload_case
    source = tmp_path / "report.pdf"
    source.write_bytes(b"VERSION A")
    if existing:
        (uploads / source.name).write_bytes(b"OLD UPLOAD")
    conversion_paths = []
    changes = []

    def wrap_copy(real_copy):
        def copy_then_change(base, name, src, **kwargs):
            result = real_copy(base, name, src, **kwargs)
            if Path(src) == source:
                changes.append(change)
                if change == "edit":
                    source.write_bytes(b"VERSION B")
                else:
                    source.unlink()
            return result

        return copy_then_change

    async def convert(path, output_path=None):
        conversion_paths.append(path)
        output_path.write_bytes(b"CONVERTED:" + path.read_bytes())
        return output_path

    monkeypatch.setattr(client_module, "copy_upload_file_no_symlink", wrap_copy(client_module.copy_upload_file_no_symlink))
    monkeypatch.setattr(client_module, "publish_upload_copy", wrap_copy(client_module.publish_upload_copy))
    monkeypatch.setattr(conversion, "CONVERTIBLE_EXTENSIONS", {".pdf"})
    monkeypatch.setattr(conversion, "convert_file_to_markdown", convert)
    response = call_upload(client, source, in_loop)

    assert response["success"]
    assert changes == [change]
    saved = uploads / response["files"][0]["filename"]
    assert saved.read_bytes() == b"VERSION A"
    companion = resolve_companion(saved)
    assert companion is not None
    assert companion.read_bytes() == b"CONVERTED:VERSION A"
    assert conversion_paths[0].name == source.name
    assert conversion_paths[0] != source
    assert not conversion_paths[0].is_relative_to(uploads)
    assert not conversion_paths[0].parent.exists()


@pytest.mark.parametrize("in_loop", [False, True])
def test_failed_conversion_keeps_original_and_cleans_snapshot(upload_case, monkeypatch, tmp_path, in_loop):
    import deerflow.utils.file_conversion as conversion

    client, uploads = upload_case
    source = tmp_path / "report.pdf"
    source.write_bytes(b"VERSION A")
    paths = []

    async def fail(path, output_path=None):
        paths.append(path)
        raise RuntimeError("conversion failed")

    monkeypatch.setattr(conversion, "CONVERTIBLE_EXTENSIONS", {".pdf"})
    monkeypatch.setattr(conversion, "convert_file_to_markdown", fail)
    response = call_upload(client, source, in_loop)
    assert response["success"]
    assert "markdown_file" not in response["files"][0]
    assert (uploads / "report.pdf").read_bytes() == b"VERSION A"
    assert source.read_bytes() == b"VERSION A"
    assert paths[0] != source
    assert not paths[0].parent.exists()


def test_convertible_same_file_is_rejected_without_data_loss(upload_case):
    client, uploads = upload_case
    source = uploads / "report.pdf"
    source.write_bytes(b"IMPORTANT BYTES")
    with pytest.raises(shutil.SameFileError):
        client.upload_files("snapshot-thread", [source])
    assert source.read_bytes() == b"IMPORTANT BYTES"


def test_captured_upload_preserves_original_metadata(upload_case, monkeypatch, tmp_path):
    import deerflow.utils.file_conversion as conversion

    client, uploads = upload_case
    source = tmp_path / "report.pdf"
    source.write_bytes(b"VERSION A")
    source.chmod(0o640)
    os.utime(source, ns=(1_700_000_000_000_000_000, 1_700_000_001_000_000_000))
    before = source.stat()

    async def convert(path, output_path=None):
        output_path.write_bytes(path.read_bytes())
        return output_path

    monkeypatch.setattr(conversion, "CONVERTIBLE_EXTENSIONS", {".pdf"})
    monkeypatch.setattr(conversion, "convert_file_to_markdown", convert)
    client.upload_files("snapshot-thread", [source])
    saved = (uploads / "report.pdf").stat()
    if os.chmod in os.supports_fd:
        assert stat.S_IMODE(saved.st_mode) == stat.S_IMODE(before.st_mode)
    if os.utime in os.supports_fd:
        assert saved.st_mtime_ns == before.st_mtime_ns


@pytest.mark.parametrize("in_loop", [False, True])
@pytest.mark.parametrize("existing", [False, True])
def test_real_xlsx_conversion_uses_captured_version(upload_case, monkeypatch, tmp_path, in_loop, existing):
    from openpyxl import Workbook, load_workbook

    import deerflow.client as client_module

    client, uploads = upload_case
    source = tmp_path / "report.xlsx"
    workbook = Workbook()
    workbook.active["A1"] = "VERSION A"
    workbook.save(source)
    if existing:
        (uploads / source.name).write_bytes(b"OLD UPLOAD")
    saves = []

    def wrap_copy(real_copy):
        def copy_then_save(base, name, src, **kwargs):
            result = real_copy(base, name, src, **kwargs)
            if Path(src) == source:
                saves.append(source)
                workbook.active["A1"] = "VERSION B"
                workbook.save(source)
            return result

        return copy_then_save

    monkeypatch.setattr(client_module, "copy_upload_file_no_symlink", wrap_copy(client_module.copy_upload_file_no_symlink))
    monkeypatch.setattr(client_module, "publish_upload_copy", wrap_copy(client_module.publish_upload_copy))
    response = call_upload(client, source, in_loop)
    assert response["success"]
    assert saves == [source]
    saved = uploads / response["files"][0]["filename"]
    if existing:
        assert saved.name == "report_1.xlsx"
        assert (uploads / "report.xlsx").read_bytes() == b"OLD UPLOAD"
    saved_workbook = load_workbook(saved)
    try:
        assert saved_workbook.active["A1"].value == "VERSION A"
    finally:
        saved_workbook.close()
        workbook.close()
    companion = resolve_companion(saved)
    assert companion is not None
    text = companion.read_text(encoding="utf-8")
    assert "VERSION A" in text
    assert "VERSION B" not in text


def test_capture_failure_preserves_existing_upload(upload_case, monkeypatch, tmp_path):
    import deerflow.uploads.manager as manager

    client, uploads = upload_case
    source = tmp_path / "report.pdf"
    source.write_bytes(b"NEW CONTENT")
    existing = uploads / "report.pdf"
    existing.write_bytes(b"EXISTING CONTENT")
    captures = []
    real_copyfileobj = shutil.copyfileobj

    def fail_capture(src, dest, *args, **kwargs):
        src_name, dest_name = getattr(src, "name", None), getattr(dest, "name", None)
        if isinstance(src_name, (str, os.PathLike)) and isinstance(dest_name, (str, os.PathLike)):
            capture = Path(dest_name)
            if Path(src_name) == source and capture.parent.name.startswith("deerflow-upload-source-"):
                captures.append(capture)
                raise OSError("snapshot disk full")
        return real_copyfileobj(src, dest, *args, **kwargs)

    monkeypatch.setattr(manager.shutil, "copyfileobj", fail_capture)
    unrelated_copy = BytesIO()
    shutil.copyfileobj(BytesIO(b"UNRELATED CONTENT"), unrelated_copy)
    assert unrelated_copy.getvalue() == b"UNRELATED CONTENT"
    with pytest.raises(OSError, match="snapshot disk full"):
        client.upload_files("snapshot-thread", [source])
    assert existing.read_bytes() == b"EXISTING CONTENT"
    assert source.read_bytes() == b"NEW CONTENT"
    assert not captures[0].parent.exists()


@pytest.mark.skipif(sys.platform != "darwin", reason="chmod +a ACL grants require macOS")
@pytest.mark.parametrize("in_loop", [False, True])
def test_acl_only_readable_source_uploads_and_converts(upload_case, monkeypatch, tmp_path, in_loop):
    """Preserve the ACL-only source case identified in the #6621 review."""
    import deerflow.utils.file_conversion as conversion

    if os.geteuid() == 0:
        pytest.skip("root bypasses source permission checks")
    client, uploads = upload_case
    source = tmp_path / "report.pdf"
    source.write_bytes(b"ACL-AUTHORIZED CONTENT")
    source.chmod(0o000)
    saved = uploads / source.name
    conversion_paths = []

    async def convert(path, output_path=None):
        conversion_paths.append(path)
        output_path.write_bytes(b"CONVERTED:" + path.read_bytes())
        return output_path

    monkeypatch.setattr(conversion, "CONVERTIBLE_EXTENSIONS", {".pdf"})
    monkeypatch.setattr(conversion, "convert_file_to_markdown", convert)
    try:
        subprocess.run(["chmod", "+a", f"user:{getpass.getuser()} allow read", str(source)], check=True)
        assert stat.S_IMODE(source.stat().st_mode) == 0
        assert source.read_bytes() == b"ACL-AUTHORIZED CONTENT"
        response = call_upload(client, source, in_loop)
        assert response["success"]
        companion = resolve_companion(saved)
        assert companion is not None
        assert companion.read_bytes() == b"CONVERTED:ACL-AUTHORIZED CONTENT"
        if os.chmod in os.supports_fd:
            assert stat.S_IMODE(saved.stat().st_mode) == 0
        # The uploaded copy preserves mode bits, but not the source ACL.
        saved.chmod(0o600)
        assert saved.read_bytes() == b"ACL-AUTHORIZED CONTENT"
        assert conversion_paths[0] != source
        assert not conversion_paths[0].is_relative_to(uploads)
        assert not conversion_paths[0].parent.exists()
    finally:
        source.chmod(0o600)
        if saved.exists():
            saved.chmod(0o600)


@pytest.mark.parametrize("in_loop", [False, True])
@pytest.mark.parametrize("change", ["edit", "remove"])
def test_collision_retries_keep_one_capture(upload_case, monkeypatch, tmp_path, in_loop, change):
    """Retry new names without reopening the source or losing its metadata."""
    import builtins

    import deerflow.uploads.manager as manager
    import deerflow.utils.file_conversion as conversion

    client, uploads = upload_case
    source = tmp_path / "report.pdf"
    source.write_bytes(b"VERSION A")
    source.chmod(0o640)
    os.utime(source, ns=(1_700_000_000_000_000_000, 1_700_000_001_000_000_000))
    before = source.stat()
    real_open = builtins.open
    real_link = manager.os.link
    source_opens = []
    attempts = []
    captures = []

    def record_open(path, mode="r", *args, **kwargs):
        if mode == "rb" and isinstance(path, (str, os.PathLike)) and Path(path) == source:
            source_opens.append(path)
        return real_open(path, mode, *args, **kwargs)

    def claim_then_link(staging, destination, **kwargs):
        if destination.suffix == ".pdf":
            attempts.append(destination.name)
            if len(attempts) == 1:
                # Another caller wins two names while this call is publishing.
                (uploads / "report.pdf").write_bytes(b"FIRST OWNER")
                (uploads / "report_1.pdf").write_bytes(b"SECOND OWNER")
                if change == "edit":
                    source.write_bytes(b"VERSION B")
                    source.chmod(0o600)
                    os.utime(source, ns=(1_800_000_000_000_000_000, 1_800_000_001_000_000_000))
                else:
                    source.unlink()
        return real_link(staging, destination, **kwargs)

    async def convert(path, output_path=None):
        captures.append(path)
        output_path.write_bytes(b"CONVERTED:" + path.read_bytes())
        return output_path

    monkeypatch.setattr(builtins, "open", record_open)
    monkeypatch.setattr(manager.os, "link", claim_then_link)
    monkeypatch.setattr(conversion, "CONVERTIBLE_EXTENSIONS", {".pdf"})
    monkeypatch.setattr(conversion, "convert_file_to_markdown", convert)
    response = call_upload(client, source, in_loop)

    assert response["success"]
    info = response["files"][0]
    assert info["filename"] == "report_2.pdf"
    assert info["original_filename"] == "report.pdf"
    assert attempts == ["report.pdf", "report_1.pdf", "report_2.pdf"]
    assert len(source_opens) == 1
    saved = uploads / info["filename"]
    assert saved.read_bytes() == b"VERSION A"
    companion = resolve_companion(saved)
    assert companion == uploads / "report_2.md"
    assert companion.read_bytes() == b"CONVERTED:VERSION A"
    assert (uploads / "report.pdf").read_bytes() == b"FIRST OWNER"
    assert (uploads / "report_1.pdf").read_bytes() == b"SECOND OWNER"
    if os.chmod in os.supports_fd:
        assert stat.S_IMODE(saved.stat().st_mode) == stat.S_IMODE(before.st_mode)
    if os.utime in os.supports_fd:
        assert saved.stat().st_mtime_ns == before.st_mtime_ns
    assert captures[0] != source
    assert not captures[0].is_relative_to(uploads)
    assert not captures[0].parent.exists()
    assert not list(uploads.glob(".upload-*.part"))


def test_colliding_original_is_published_only_after_copy_completes(upload_case, monkeypatch, tmp_path):
    import deerflow.uploads.manager as manager

    client, uploads = upload_case
    source = tmp_path / "notes.txt"
    source.write_bytes(b"NEW UPLOAD")
    (uploads / source.name).write_bytes(b"OLD UPLOAD")
    observed = []
    real_copy = manager.shutil.copyfileobj

    def inspect_copy(src, dst, *args, **kwargs):
        real_copy(src, dst, *args, **kwargs)
        observed.append({item["filename"] for item in manager.list_files_in_dir(uploads)["files"]})

    monkeypatch.setattr(manager.shutil, "copyfileobj", inspect_copy)
    response = client.upload_files("snapshot-thread", [source])

    assert observed
    assert all(names == {"notes.txt"} for names in observed)
    assert response["files"][0]["filename"] == "notes_1.txt"
    assert (uploads / "notes.txt").read_bytes() == b"OLD UPLOAD"
    assert (uploads / "notes_1.txt").read_bytes() == b"NEW UPLOAD"
