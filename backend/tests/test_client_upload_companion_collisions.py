"""Converted uploads must preserve Markdown from earlier upload calls."""

import asyncio
from pathlib import Path

import pytest
from openpyxl import Workbook
from pptx import Presentation

from deerflow.client import DeerFlowClient
from deerflow.uploads.companions import resolve_companion


def _document(path: Path, content: str) -> None:
    if path.suffix == ".md":
        path.write_text(content, encoding="utf-8")
    elif path.suffix == ".xlsx":
        book = Workbook()
        book.active["A1"] = content
        book.save(path)
        book.close()
    else:
        slides = Presentation()
        slide = slides.slides.add_slide(slides.slide_layouts[1])
        slide.shapes.title.text = content
        slides.save(path)


@pytest.fixture
def upload_case(tmp_path, monkeypatch):
    import deerflow.client as client_module

    uploads = tmp_path / "thread/user-data/uploads"
    uploads.mkdir(parents=True)
    monkeypatch.setattr(client_module, "ensure_uploads_dir", lambda *args, **kwargs: uploads)
    return DeerFlowClient.__new__(DeerFlowClient), uploads


@pytest.mark.parametrize("in_loop", [False, True])
@pytest.mark.parametrize("first_ext,second_ext", [(".md", ".xlsx"), (".md", ".pptx"), (".xlsx", ".pptx"), (".pptx", ".xlsx")])
def test_separate_uploads_preserve_existing_markdown(upload_case, tmp_path, in_loop, first_ext, second_ext):
    client, uploads = upload_case
    first = tmp_path / f"report{first_ext}"
    second = tmp_path / f"report{second_ext}"
    _document(first, "KEEP FIRST DOCUMENT")
    _document(second, "SECOND DOCUMENT")

    def upload(source):
        if not in_loop:
            return client.upload_files("companion-thread", [source])

        # Run the synchronous upload inside an active event loop to exercise conversion_pool.
        async def run():
            return client.upload_files("companion-thread", [source])

        return asyncio.run(run())

    first_response = upload(first)
    original_markdown = uploads / "report.md"
    original_bytes = original_markdown.read_bytes()
    second_response = upload(second)

    assert first_response["success"] and second_response["success"]
    assert original_markdown.read_bytes() == original_bytes
    assert "KEEP FIRST DOCUMENT" in original_markdown.read_text(encoding="utf-8")
    info = second_response["files"][0]
    assert info["filename"] == second.name
    assert info["markdown_file"] == "report_1.md"
    companion = uploads / info["markdown_file"]
    assert "SECOND DOCUMENT" in companion.read_text(encoding="utf-8")
    assert resolve_companion(uploads / second.name) == companion
    assert Path(info["markdown_path"]) == companion
    assert info["markdown_virtual_path"].endswith("/report_1.md")
    assert info["markdown_artifact_url"].endswith("/report_1.md")
    assert (uploads / first.name).read_bytes() == first.read_bytes()
    assert (uploads / second.name).read_bytes() == second.read_bytes()
    if first_ext != ".md":
        assert resolve_companion(uploads / first.name) == original_markdown


def test_reupload_replaces_original_but_preserves_occupied_companion_names(upload_case, tmp_path):
    client, uploads = upload_case
    source = tmp_path / "report.xlsx"
    _document(source, "FIRST DOCUMENT")
    client.upload_files("companion-thread", [source])
    first_markdown = (uploads / "report.md").read_bytes()
    occupied = uploads / "report_1.md"
    occupied.write_text("INDEPENDENT NOTES", encoding="utf-8")

    _document(source, "REPLACEMENT DOCUMENT")
    response = client.upload_files("companion-thread", [source])

    assert response["success"]
    info = response["files"][0]
    assert info["filename"] == "report.xlsx"
    assert (uploads / source.name).read_bytes() == source.read_bytes()
    assert info["markdown_file"] == "report_2.md"
    assert (uploads / "report.md").read_bytes() == first_markdown
    assert occupied.read_text(encoding="utf-8") == "INDEPENDENT NOTES"
    companion = uploads / "report_2.md"
    assert "REPLACEMENT DOCUMENT" in companion.read_text(encoding="utf-8")
    assert resolve_companion(uploads / source.name) == companion
